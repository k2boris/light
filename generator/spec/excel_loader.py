"""Excel loader for V4.1 workbooks.

We use openpyxl to read Excel deterministically and without formula evaluation.
All values are read as displayed cell values (strings/numbers) and normalized.

Important: row_index refers to the Excel row number (1-based), including header row.
We typically store data row indices as the row number in the worksheet for user-friendly errors.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl

from .spec_model import (
    WorkbookSpec, EntityPlan, FieldMap, EAVRule, DictEntry, Override,
    TemplateDef, TemplateParam, PatternDef, BroadwayActor, TemplateImplementation, MultiRowRule
)

logger = logging.getLogger(__name__)

REQUIRED_TABS = [
    "README","EntityPlans","FieldMaps","EAVRules","Dictionaries","Overrides",
    "Templates","Template_Params","Patterns","BroadwayActors","TemplateImplementations"
]

def _norm(v: Any) -> Optional[str]:
    """Normalize a cell value to a stripped string or None."""
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s if s != "" else None
    return str(v).strip()

def _read_table(ws, header_row: int = 1) -> Tuple[List[str], List[Dict[str, Any]]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    if header_row < 1:
        header_row = 1
    if header_row > len(rows):
        return [], []
    header_idx = header_row - 1
    header = [str(h).strip() if h is not None else "" for h in rows[header_idx]]
    out = []
    for i, r in enumerate(rows[header_idx + 1 :], start=header_row + 1):  # excel row number
        rec = {}
        empty = True
        for j, col in enumerate(header):
            if col == "":
                continue
            val = r[j] if j < len(r) else None
            n = _norm(val)
            if n is not None:
                empty = False
            rec[col] = n
        if not empty:
            rec["__row_index__"] = i
            out.append(rec)
    return header, out

def load_v4_1_workbook(path: Path) -> WorkbookSpec:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    for tab in REQUIRED_TABS:
        if tab not in wb.sheetnames:
            raise ValueError(f"Missing required tab '{tab}' in workbook {path}")

    spec = WorkbookSpec()

    # EntityPlans
    ws = wb["EntityPlans"]
    entity_rows = list(ws.iter_rows(values_only=True))
    header_row = 1
    if entity_rows:
        row1 = [str(v).strip() for v in entity_rows[0] if v is not None]
        if "plan_id" not in row1 and len(entity_rows) >= 2:
            row2 = [str(v).strip() for v in entity_rows[1] if v is not None]
            if "plan_id" in row2:
                header_row = 2
    _, rows = _read_table(ws, header_row=header_row)
    for rec in rows:
        target_pk_cols = rec.get("target_pk_cols") or rec.get("keys") or ""
        input_port = rec.get("input_port") or rec.get("driver_key") or ""
        input_args = rec.get("input_args") or rec.get("driver_args")
        ep = EntityPlan(
            row_index=rec["__row_index__"],
            plan_id=rec.get("plan_id") or "",
            target_table=rec.get("target_table") or "",
            target_entity=rec.get("target_entity"),
            target_grain=rec.get("target_grain") or "",
            keys=target_pk_cols,
            source_system=rec.get("source_system") or "",
            source_table=rec.get("source_table") or "",
            source_entity=rec.get("source_entity") or "",
            template_id=rec.get("template_id") or "",
            driver_key=input_port,
            driver_args=input_args,
            row_policy=rec.get("row_policy") or "",
            multiplicity_expectation=rec.get("multiplicity_expectation") or "",
            as_of_policy=rec.get("as_of_policy") or "",
            source_interface=rec.get("source_interface"),
            target_interface=rec.get("target_interface"),
            target_schema=rec.get("target_schema"),
            dialect=rec.get("dialect"),
            iterate_mode=rec.get("iterate_mode"),
            notes=rec.get("notes"),
        )
        spec.entity_plans[ep.plan_id] = ep

    # FieldMaps
    ws = wb["FieldMaps"]
    _, rows = _read_table(ws)
    for rec in rows:
        fm = FieldMap(
            row_index=rec["__row_index__"],
            plan_id=rec.get("plan_id") or "",
            target_table=rec.get("target_table") or "",
            target_field=rec.get("target_field") or "",
            target_description=rec.get("target_description"),
            source_path=rec.get("source_path"),
            pattern_id=rec.get("pattern_id") or "",
            pattern_args=rec.get("pattern_args"),
            required=rec.get("required") or "N",
            default=rec.get("default"),
            transform_rule=rec.get("transform_rule"),
            join_key=rec.get("join_key"),
            notes=rec.get("notes"),
        )
        spec.field_maps.append(fm)
        spec.field_maps_by_plan.setdefault(fm.plan_id, []).append(fm)

    # EAVRules
    ws = wb["EAVRules"]
    _, rows = _read_table(ws)
    for rec in rows:
        er = EAVRule(
            row_index=rec["__row_index__"],
            plan_id=rec.get("plan_id") or "",
            target_table=rec.get("target_table") or "",
            driver_key=rec.get("driver_key") or "",
            name=rec.get("name"),
            name_expr=rec.get("name_expr"),
            value_source_path=rec.get("value_source_path"),
            value_expr=rec.get("value_expr"),
            value_type=rec.get("value_type") or "",
            join_key=rec.get("join_key"),
            notes=rec.get("notes"),
            transform_rule=rec.get("transform_rule"),
        )
        spec.eav_rules.append(er)
        spec.eav_rules_by_plan.setdefault(er.plan_id, []).append(er)

    # MultiRowRules (optional, used by multi-row templates)
    if "MultiRowRules" in wb.sheetnames:
        ws = wb["MultiRowRules"]
        _, rows = _read_table(ws)
        for rec in rows:
            mr = MultiRowRule(
                row_index=rec["__row_index__"],
                plan_id=rec.get("plan_id") or "",
                row_id=rec.get("row_id") or "",
                emit_when=rec.get("emit_when"),
                target_field=rec.get("target_field") or "",
                pattern_id=rec.get("pattern_id") or "",
                source_path=rec.get("source_path"),
                pattern_args=rec.get("pattern_args"),
                required=rec.get("required") or "N",
                notes=rec.get("notes"),
            )
            spec.multirow_rules.append(mr)
            spec.multirow_rules_by_plan.setdefault(mr.plan_id, []).append(mr)

    # Dictionaries
    ws = wb["Dictionaries"]
    _, rows = _read_table(ws)
    for rec in rows:
        de = DictEntry(
            row_index=rec["__row_index__"],
            dict_type=rec.get("dict_type") or "",
            name=rec.get("name") or "",
            key=rec.get("key") or "",
            value=rec.get("value") or "",
            notes=rec.get("notes"),
        )
        spec.dictionaries.append(de)

    # Overrides
    ws = wb["Overrides"]
    _, rows = _read_table(ws)
    for rec in rows:
        ov = Override(
            row_index=rec["__row_index__"],
            plan_id=rec.get("plan_id"),
            target_table=rec.get("target_table"),
            override_type=rec.get("override_type") or "",
            name=rec.get("name") or "",
            params=rec.get("params") or "{}",
            notes=rec.get("notes"),
        )
        spec.overrides.append(ov)

    # Templates
    ws = wb["Templates"]
    _, rows = _read_table(ws)
    for rec in rows:
        td = TemplateDef(
            row_index=rec["__row_index__"],
            template_id=rec.get("template_id") or "",
            version=rec.get("version") or "",
            description=rec.get("description") or "",
            ir_skeleton_json=rec.get("ir_skeleton_json") or "{}",
            notes=rec.get("notes"),
        )
        spec.templates[td.template_id] = td

    # Template_Params
    ws = wb["Template_Params"]
    _, rows = _read_table(ws)
    for rec in rows:
        tp = TemplateParam(
            row_index=rec["__row_index__"],
            template_id=rec.get("template_id") or "",
            param_name=rec.get("param_name") or "",
            required=rec.get("required") or "N",
            type=rec.get("type") or "string",
            default=rec.get("default"),
            source=rec.get("source") or "",
            notes=rec.get("notes"),
        )
        spec.template_params.append(tp)

    # Patterns
    ws = wb["Patterns"]
    _, rows = _read_table(ws)
    for rec in rows:
        pd = PatternDef(
            row_index=rec["__row_index__"],
            pattern_id=rec.get("pattern_id") or "",
            version=rec.get("version") or "",
            description=rec.get("description") or "",
            expr_template=rec.get("expr_template") or "",
            required_inputs=rec.get("required_inputs") or "",
            notes=rec.get("notes"),
        )
        spec.patterns[pd.pattern_id] = pd

    # BroadwayActors
    ws = wb["BroadwayActors"]
    _, rows = _read_table(ws)
    for rec in rows:
        ba = BroadwayActor(
            row_index=rec["__row_index__"],
            template_id=rec.get("template_id") or "",
            stage=rec.get("stage") or "",
            actor_id=rec.get("actor_id") or "",
            parent=rec.get("parent") or "",
            enabled=rec.get("enabled") or "Y",
            in_json=rec.get("in_json"),
            out_json=rec.get("out_json"),
            notes=rec.get("notes"),
        )
        spec.broadway_actors.append(ba)
        spec.broadway_actors_by_template.setdefault(ba.template_id, []).append(ba)

    # TemplateImplementations
    ws = wb["TemplateImplementations"]
    _, rows = _read_table(ws)
    for rec in rows:
        ti = TemplateImplementation(
            row_index=rec["__row_index__"],
            template_id=rec.get("template_id") or "",
            runtime=(rec.get("runtime") or "").lower(),
            enabled=rec.get("enabled") or "Y",
            impl_json=rec.get("impl_json"),
            notes=rec.get("notes"),
        )
        spec.template_implementations.append(ti)
        spec.template_impls_by_template.setdefault(ti.template_id, {})[ti.runtime] = ti

    logger.info(
        f"Loaded spec: plans={len(spec.entity_plans)} field_maps={len(spec.field_maps)} eav_rules={len(spec.eav_rules)} multirow_rules={len(spec.multirow_rules)} broadway_actors={len(spec.broadway_actors)} template_impls={len(spec.template_implementations)}",
        extra={"phase":"SPEC"},
    )
    return spec
